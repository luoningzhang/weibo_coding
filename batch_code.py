"""
批量编码微博（云雾 API / claude-sonnet-4-6）

用法（原有模式）:
    python batch_code.py --input data/xxx.json --output data/out.json

用法（按电影批量模式）:
    python batch_code.py --movies 流浪地球2 封神 长安三万里
    python batch_code.py --movies all          # 处理 data/tmp/ 下所有 *_电影 文件夹

可选参数:
    --codebook  data/codebook.xlsx
    --batch     20      每批条数（默认 20）
    --workers   5       并发数
    --tmp-dir   data/tmp   电影文件夹根目录
    --show-prompt       打印 system prompt 后退出
    --test              跑前200条与 top200-coding.xlsx 对比后退出

API Key 放在 config.json 的 api_key 字段中。
"""

import argparse
import json
import random
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

from openpyxl import load_workbook
from openai import OpenAI

YUNWU_BASE_URL  = "https://yunwu.ai/v1"
MODEL           = "claude-sonnet-4-6-thinking"
MAX_RETRIES     = 3
MAX_TOKENS      = 6000
THINKING_BUDGET = 4000
CONFIG_PATH     = Path(__file__).parent / "config.json"


# ── 特殊异常：Token 耗尽，需立即停止 ─────────────────────────────
class QuotaExhaustedError(Exception):
    pass


# ── 进度条 ────────────────────────────────────────────────────────
_progress_bar: "ProgressBar | None" = None


def log(msg: str):
    """打印一行信息，不破坏进度条位置。"""
    if _progress_bar is not None:
        with _progress_bar._lock:
            # \n 先换行离开进度条行，不擦除它
            # 下次 update() 的 \r 会在新当前行重绘进度条
            sys.stderr.write(f"\n{msg}\n")
            sys.stderr.flush()
    else:
        sys.stderr.write(msg + "\n")
        sys.stderr.flush()


class ProgressBar:
    BAR_WIDTH = 40

    def __init__(self, total: int):
        global _progress_bar
        self.total  = total
        self.done   = 0
        self._lock  = threading.Lock()
        self._start = time.time()
        _progress_bar = self

    def update(self, n: int = 1):
        with self._lock:
            self.done += n
            self._render()

    def _render(self):
        pct     = self.done / self.total if self.total else 0
        filled  = int(self.BAR_WIDTH * pct)
        bar     = "█" * filled + "░" * (self.BAR_WIDTH - filled)
        elapsed = time.time() - self._start
        speed   = self.done / elapsed if elapsed > 0 else 0
        remain  = (self.total - self.done) / speed if speed > 0 else 0
        sys.stderr.write(
            f"\r[{bar}] {pct*100:.1f}% | {self.done}/{self.total} | "
            f"{speed:.1f}条/秒 | 已用{elapsed/60:.1f}分 | 剩余{remain/60:.1f}分  "
        )
        sys.stderr.flush()

    def finish(self):
        global _progress_bar
        with self._lock:
            self._render()
        sys.stderr.write("\n")
        sys.stderr.flush()
        _progress_bar = None


# ── 工具函数 ──────────────────────────────────────────────────────
def load_config() -> dict:
    with open(CONFIG_PATH, encoding="utf-8") as f:
        return json.load(f)


def load_codebook(path: str) -> list[dict]:
    wb = load_workbook(path)
    ws = wb.active
    codes = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        num, name, definition, criteria, note = (list(row) + [None] * 5)[:5]
        if num is None:
            continue
        codes.append({
            "num":        int(num),
            "name":       str(name).strip(),
            "definition": str(definition).strip() if definition else "",
            "criteria":   str(criteria).strip() if criteria else "",
            "note":       str(note).strip() if note else "",
        })
    return codes


def load_top200_coding(path: str) -> dict[str, list[int]]:
    import re
    wb = load_workbook(path)
    ws = wb.active
    result = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        weibo_id = str(row[1]).strip() if row[1] is not None else None
        coding_str = row[12]
        if not weibo_id:
            continue
        nums = [int(n) for n in re.findall(r'\b(\d+)\b', str(coding_str))] if coding_str else []
        result[weibo_id] = nums
    return result


def build_system_prompt(codes: list[dict], fewshot_path: str = None) -> str:
    lines = [
        "你是一位内容分析员，负责对微博文本按行为类别编码。",
        "每条微博可对应 0 到多个编码。只返回 JSON，不要任何解释。",
        "",
        "## 编码列表",
    ]
    for c in codes:
        lines.append(f"\n### {c['num']} {c['name']}")
        lines.append(f"操作定义：{c['definition']}")
        lines.append(f"识别标准：{c['criteria']}")
        if c["note"]:
            lines.append(f"备注：{c['note']}")
    lines += [
        "",
        "## 编码规则与常见易错",
        "",
        "### 开始编码前，先扫描以下高频信号",
        "在对每条微博开始编码之前，先逐一检查以下信号是否存在：",
        '- 【11 安利】正文含"快去看/强推/一定要去/值得一看/安利/带朋友去看"等 → 必须加11',
        '- 【17 自来水】正文含"自来水/精神股东/不是粉丝/路人" → 必须加17',
        "- 【19 共情】正文对导演/演员/幕后有心疼/感动/感谢/敬佩 → 必须加19",
        "- 【22 市场分析】正文同时提及两部或以上影片并进行比较 → 必须加22",
        "",
        "### 各代码精确边界",
        "",
        "**代码8（长评）**：必须是100字以上、有明确论点导向的分析。以下不算：纯情绪感叹、剧情复述、粉丝向角色讨论、列优缺点但无论点。",
        "",
        "**代码12（话题词条）**：使用#话题#标签是微博常规操作，不构成此代码。必须有明确的词条聚合引导意图（如呼吁他人到该话题下跟帖、在超话内组织讨论）才算。",
        "",
        "**代码14 vs 15 vs 22**：",
        "- 有具体批评论点且逐条驳斥，或质疑批评者动机/来源 → 14",
        "- 无具体批评对象，为本片理性澄清立场 → 15",
        "- 涉及不同影片之间的比较（票房/质量/排片） → 22",
        "- 三者可以同时出现",
        "",
        "**代码5 vs 7**：",
        "- 有创作成本（绘图/剪辑/同人文/音频） → 5；发了图/视频链接即打5",
        "- 纯文字谐音梗/改编段子/截图加文字调侃 → 7",
        "",
        "**代码11备注**：必须有推荐动作；带朋友去看、叫朋友一起也算；仅@某人讨论不算。",
        "",
        "**代码6备注**：是作为观众主动发现细节，不是作为宣发/制作方披露。",
        "",
        '**空数组规则**：只有以下情况才允许输出空数组[]：正文是纯转发无评论、纯图片说明、或与电影完全无关的内容。只要正文有实质性的涉片内容，就必须至少编一个代码。禁止把"拿不准"当作不编码的理由。',
        "",
    ]
    if fewshot_path and Path(fewshot_path).exists():
        with open(fewshot_path, encoding="utf-8") as f:
            lines.append(f.read())
    lines += [
        "## 输出格式",
        "输入是一个 JSON 数组，每项含 id/user/movie/text 字段，其中 id 是序号（1开始）。",
        "输出必须是一个 JSON 对象，key 是 id（字符串），value 是该微博适用的编码编号数组。",
        "没有适用编码则为空数组[]。必须为每一条都输出对应的 key。",
        '例：输入3条(id=1,2,3) → 输出 {"1":[5,12],"2":[22],"3":[]}',
    ]
    return "\n".join(lines)


def load_posts(path: str) -> list[dict]:
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    if isinstance(data, list):
        return data
    for key in ("data", "posts", "statuses", "items"):
        if key in data and isinstance(data[key], list):
            return data[key]
    raise ValueError(f"无法识别 JSON 结构: {type(data)}")


def post_id(post: dict) -> str:
    return str(post.get("id") or post.get("mid") or post.get("idstr", ""))


def is_quota_error(err_str: str) -> bool:
    keywords = ["insufficient_quota", "quota exceeded", "402", "billing",
                "credit", "out of token", "余额不足", "账户余额"]
    return any(k in err_str.lower() for k in keywords)


def is_content_filter(err_str: str) -> bool:
    return "1301" in err_str or "contentfilter" in err_str.lower()


def _extract_result_json(raw: str, batch_index: int) -> dict:
    """从模型原始响应中提取编码结果 dict，跳过 thinking 内容里的 {...}。"""
    decoder = json.JSONDecoder()
    pos = 0
    while pos < len(raw):
        start = raw.find("{", pos)
        if start == -1:
            break
        try:
            obj, _ = decoder.raw_decode(raw, start)
            # 有效结果：dict 且所有 key 都是纯数字字符串
            if isinstance(obj, dict) and obj and all(k.isdigit() for k in obj):
                return obj
        except json.JSONDecodeError:
            pass
        pos = start + 1
    raise ValueError(f"[批次{batch_index}] 响应中未找到编码结果 JSON: {raw[:300]}")


def call_api(client: OpenAI, system_prompt: str,
             batch: list[dict], batch_index: int) -> list[list[int]]:
    payload = [
        {
            "id": i + 1,
            "user": (p.get("user", {}).get("screen_name", "")
                     if isinstance(p.get("user"), dict)
                     else str(p.get("user", ""))),
            "movie": p.get("movie", ""),
            "text": (p.get("text") or p.get("content") or "")[:500],
        }
        for i, p in enumerate(batch)
    ]
    user_msg  = json.dumps(payload, ensure_ascii=False)
    user_msg += f"\n<!-- run_id:{random.randint(10000,99999)} -->"

    for attempt in range(1, MAX_RETRIES + 1):
        try:
            resp = client.chat.completions.create(
                model=MODEL,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user",   "content": user_msg},
                ],
                max_tokens=MAX_TOKENS,
                extra_body={"thinking": {"type": "enabled",
                                         "budget_tokens": THINKING_BUDGET}},
            )
            raw   = resp.choices[0].message.content.strip()
            result_dict = _extract_result_json(raw, batch_index)
            result, missing = [], []
            for i in range(len(batch)):
                key = str(i + 1)
                if key in result_dict:
                    result.append(result_dict[key])
                else:
                    result.append([])
                    missing.append(key)
            if missing:
                log(f"  [批次{batch_index}] 警告：缺失 id={','.join(missing)}，已填空")
            return result

        except Exception as e:
            err_str = str(e)

            # ① 内容过滤：直接跳过，不重试
            if is_content_filter(err_str):
                log(f"⚠️  [内容过滤] 批次{batch_index} 触发内容过滤，填充空编码跳过")
                return [[] for _ in batch]

            # ② Token/余额耗尽：向上抛出，让主流程停止
            if is_quota_error(err_str):
                raise QuotaExhaustedError(f"API 余额/Token 耗尽：{err_str}")

            # ③ 其他错误：正常重试
            wait = 2 ** attempt
            log(f"  [批次{batch_index}] 第{attempt}/{MAX_RETRIES}次失败: {err_str[:120]}")
            if attempt < MAX_RETRIES:
                log(f"  等待 {wait}s 后重试…")
                time.sleep(wait)
            else:
                log(f"  [批次{batch_index}] 放弃，填充空编码")
                return [[] for _ in batch]


# ── 核心处理函数（一个电影/文件） ────────────────────────────────
def process_file(client, system_prompt, code_name_map,
                 input_path: Path, output_path: Path,
                 batch_size: int, workers: int):

    checkpoint_path = output_path.with_suffix(".checkpoint.json")

    log(f"\n{'='*60}")
    log(f"  输入: {input_path}")
    log(f"  输出: {output_path}")

    posts = load_posts(str(input_path))
    log(f"  共 {len(posts)} 条")

    # 断点续跑
    if checkpoint_path.exists():
        with open(checkpoint_path, encoding="utf-8") as f:
            code_results: dict[str, list[int]] = json.load(f)
        log(f"  断点续跑：已完成 {len(code_results)} 条")
    else:
        code_results = {}

    todo = [p for p in posts if post_id(p) not in code_results]
    if not todo:
        log("  全部已完成，跳过。")
        _write_output(posts, code_results, code_name_map, output_path)
        return

    log(f"  待处理 {len(todo)} 条，批大小={batch_size}，并发={workers}")

    batches  = [todo[i:i + batch_size] for i in range(0, len(todo), batch_size)]
    progress = ProgressBar(len(todo))
    quota_hit = False

    def process_batch(idx_batch):
        idx, b = idx_batch
        return b, call_api(client, system_prompt, b, idx)

    try:
        with ThreadPoolExecutor(max_workers=workers) as executor:
            futures = {executor.submit(process_batch, (i, b)): i
                       for i, b in enumerate(batches)}
            for future in as_completed(futures):
                try:
                    b, res = future.result()
                except QuotaExhaustedError as qe:
                    log(f"🚨 [Token耗尽] {qe}")
                    log("   已保存当前进度，请充值后续跑。")
                    # 取消所有未完成 future
                    for f in futures:
                        f.cancel()
                    quota_hit = True
                    break

                for p, nums in zip(b, res):
                    code_results[post_id(p)] = [int(n) for n in nums
                                                if isinstance(n, (int, float))]
                progress.update(len(b))

                # 每批存一次 checkpoint
                with open(checkpoint_path, "w", encoding="utf-8") as f:
                    json.dump(code_results, f, ensure_ascii=False)

    except KeyboardInterrupt:
        log("⚠️  手动中断，已保存进度。")
        quota_hit = True  # 跳过最终输出写入

    progress.finish()

    if quota_hit:
        log(f"  进度已保存至 {checkpoint_path}，完成 {len(code_results)}/{len(posts)} 条。")
        return

    _write_output(posts, code_results, code_name_map, output_path)
    log(f"  ✅ 完成，已写出 → {output_path}")


def _write_output(posts, code_results, code_name_map, output_path: Path):
    output = []
    for p in posts:
        pid  = post_id(p)
        nums = code_results.get(pid, [])
        output.append({
            **p,
            "codes":      nums,
            "code_names": [f"{n} {code_name_map.get(n, '')}" for n in nums],
        })
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)


# ── CLI ───────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(description="批量编码微博")
    parser.add_argument("--input",    default="data/no_verified_classified_with_movie.json",
                        help="原有单文件模式输入")
    parser.add_argument("--output",   default="data/coded_output.json",
                        help="原有单文件模式输出")
    parser.add_argument("--codebook", default="data/codebook.xlsx")
    parser.add_argument("--batch",    type=int, default=20)
    parser.add_argument("--workers",  type=int, default=5)
    parser.add_argument("--tmp-dir",  default="data/tmp",
                        help="电影文件夹根目录（--movies 模式用）")
    parser.add_argument("--movies",   nargs="+",
                        help='指定电影名称（不含"_电影"后缀），或 all 处理全部')
    parser.add_argument("--show-prompt", action="store_true")
    parser.add_argument("--test",        action="store_true")
    parser.add_argument("--top200",   default="data/top200-coding.xlsx")
    args = parser.parse_args()

    config        = load_config()
    client        = OpenAI(api_key=config["api_key"], base_url=YUNWU_BASE_URL)
    codes         = load_codebook(args.codebook)
    system_prompt = build_system_prompt(codes, fewshot_path="data/fewshot_examples.md")
    code_name_map = {c["num"]: c["name"] for c in codes}

    if args.show_prompt:
        print(system_prompt)
        return

    # ── --movies 模式 ─────────────────────────────────────────────
    if args.movies:
        tmp_dir = Path(args.tmp_dir)

        if args.movies == ["all"]:
            folders = sorted(tmp_dir.glob("*_电影"))
            movie_names = [f.name.removesuffix("_电影") for f in folders
                           if f.is_dir()]
        else:
            movie_names = args.movies

        if not movie_names:
            log(f"❌ 在 {tmp_dir} 下找不到任何 *_电影 文件夹")
            return

        log(f"待处理电影：{movie_names}")

        for movie in movie_names:
            folder     = tmp_dir / f"{movie}_电影"
            input_path = folder / f"{movie}_电影_classified_F.json"
            output_path = folder / f"{movie}_电影_coded.json"

            if not input_path.exists():
                log(f"⚠️  找不到文件：{input_path}，跳过。")
                continue

            process_file(client, system_prompt, code_name_map,
                         input_path, output_path,
                         args.batch, args.workers)
        return

    # ── 原有单文件模式 ────────────────────────────────────────────
    if args.test:
        posts      = load_posts(args.input)
        test_posts = posts[:200]
        print(f"【测试模式】处理前 {len(test_posts)} 条，批大小={args.batch}，并发={args.workers}…")

        test_results: dict[str, list[int]] = {}

        def process_batch(idx_batch):
            idx, b = idx_batch
            return b, call_api(client, system_prompt, b, idx)

        batches  = [test_posts[i:i + args.batch] for i in range(0, len(test_posts), args.batch)]
        progress = ProgressBar(len(test_posts))
        with ThreadPoolExecutor(max_workers=args.workers) as executor:
            futures = {executor.submit(process_batch, (i, b)): i
                       for i, b in enumerate(batches)}
            for future in as_completed(futures):
                b, res = future.result()
                for p, nums in zip(b, res):
                    test_results[post_id(p)] = [int(n) for n in nums
                                                if isinstance(n, (int, float))]
                progress.update(len(b))
        progress.finish()

        ref = load_top200_coding(args.top200)
        hit = miss = extra = no_ref = 0
        print("\n" + "=" * 70)
        print(f"{'排名/ID':<20} {'人工编码':<28} {'模型编码':<28} 差异")
        print("=" * 70)
        for rank, p in enumerate(test_posts, 1):
            pid       = post_id(p)
            model_set = set(test_results.get(pid, []))
            if pid not in ref:
                no_ref += 1
                continue
            ref_set    = set(ref[pid])
            only_ref   = ref_set - model_set
            only_model = model_set - ref_set
            hit   += len(ref_set & model_set)
            miss  += len(only_ref)
            extra += len(only_model)
            if only_ref or only_model:
                ref_str   = ", ".join(str(n) for n in sorted(ref_set)) or "无"
                model_str = ", ".join(str(n) for n in sorted(model_set)) or "无"
                parts     = []
                if only_ref:
                    parts.append(f"漏:{','.join(str(n) for n in sorted(only_ref))}")
                if only_model:
                    parts.append(f"多:{','.join(str(n) for n in sorted(only_model))}")
                print(f"[{rank:03d}]{pid:<16} {ref_str:<28} {model_str:<28} {' | '.join(parts)}")

        total_ref = hit + miss
        print("=" * 70)
        print(f"\n【汇总】参考编码实例总数：{total_ref}，其中：")
        print(f"  命中：{hit}  ({hit/total_ref*100:.1f}%)" if total_ref else "  命中：0")
        print(f"  漏编：{miss}  ({miss/total_ref*100:.1f}%)" if total_ref else "  漏编：0")
        print(f"  多编：{extra}")
        if no_ref:
            print(f"  无参考：{no_ref} 条")
        return

    # 单文件正常跑
    process_file(client, system_prompt, code_name_map,
                 Path(args.input), Path(args.output),
                 args.batch, args.workers)


if __name__ == "__main__":
    main()
