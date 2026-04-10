"""
导出 SIP 行为编码帖子明细

每个 SIP 编码单独一个 sheet，最后一个 sheet 汇总所有 SIP 行为帖子（去重）。

SIP 编码范围：
  数据与市场劳动：2, 3, 4, 26
  策略性介入：    21, 22, 23, 24, 25, 27
  身份认同表达：  10, 17, 19
  话语防御：      14, 15, 16, 20

用法：
    python export_sip_codes.py
    python export_sip_codes.py --dir data/tmp --out sip_posts.xlsx
"""

import argparse
import json
from datetime import datetime
from pathlib import Path
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ── SIP 编码定义（顺序即 sheet 顺序）────────────────────────────
SIP_CODES = [
    (2,  "数据与市场劳动", "播报/分析票房数据"),
    (3,  "数据与市场劳动", "预测票房"),
    (4,  "数据与市场劳动", "分析排片/呼吁增加排片"),
    (26, "数据与市场劳动", "跨平台数据监控"),
    (21, "策略性介入",     "提出具体宣发建议"),
    (22, "策略性介入",     "分析竞争对手/市场环境"),
    (23, "策略性介入",     "讨论破圈策略"),
    (24, "策略性介入",     "主动联系/@院线/媒体"),
    (25, "策略性介入",     "地推活动"),
    (27, "策略性介入",     "海外/跨国宣发关注"),
    (10, "身份认同表达",   "组织/参与线下包场"),
    (17, "身份认同表达",   "自称精神股东/自来水"),
    (19, "身份认同表达",   "表达与主创共情"),
    (14, "话语防御",       "反驳差评/黑水"),
    (15, "话语防御",       "解释/维护电影"),
    (16, "话语防御",       "举报恶意言论"),
    (20, "话语防御",       "为电影逆袭/受委屈鸣不平"),
]
SIP_CODE_SET = {c for c, _, _ in SIP_CODES}

MOVIE_ORDER = ["流浪地球2", "消失的她", "封神", "长安三万里", "孤注一掷"]
FILTER_PATTERN = "今年我最爱的#微博年度电影#是"

# ── 样式 ────────────────────────────────────────────────────────
thin     = Side(style="thin")
BDR      = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill("solid", fgColor="2F5597")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
CAT_FILLS = {
    "数据与市场劳动": PatternFill("solid", fgColor="DDEBF7"),
    "策略性介入":     PatternFill("solid", fgColor="E2EFDA"),
    "身份认同表达":   PatternFill("solid", fgColor="FFF2CC"),
    "话语防御":       PatternFill("solid", fgColor="FCE4D6"),
}
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=False)
WRAP   = Alignment(horizontal="left",   vertical="top", wrap_text=True)
BODY   = Font(size=9)


def normalize_movie(raw: str) -> str:
    return raw.removesuffix("_电影").strip()


def engagement(p: dict) -> int:
    return (int(p.get("reposts_count") or 0)
            + int(p.get("comments_count") or 0)
            + int(p.get("attitudes_count") or 0))


def get_screen_name(p: dict) -> str:
    u = p.get("user")
    return u.get("screen_name", "") if isinstance(u, dict) else ""


def fmt_dt(raw: str) -> str:
    try:
        return datetime.strptime(raw, "%a %b %d %H:%M:%S +0800 %Y").strftime("%Y-%m-%d %H:%M")
    except Exception:
        return raw


def _iter_posts_from_file(path: Path):
    """
    逐条产生 post dict，尽量少占内存。
    优先用 ijson 流式解析；若未安装则 json.load 后立即遍历。
    """
    movie_from_file = path.stem.removesuffix("_coded").removesuffix("_电影")

    # ── 尝试 ijson 流式解析 ──────────────────────────────────────
    try:
        import ijson
        with open(path, "rb") as f:
            # 先试顶层数组 (prefix = "item")
            count = 0
            for post in ijson.items(f, "item"):
                if not post.get("movie"):
                    post["movie"] = movie_from_file
                yield post
                count += 1
            if count > 0:
                return
        # 顶层是 {"data":[...]} 等结构
        for prefix in ("data.item", "posts.item", "items.item"):
            with open(path, "rb") as f:
                count = 0
                try:
                    for post in ijson.items(f, prefix):
                        if not post.get("movie"):
                            post["movie"] = movie_from_file
                        yield post
                        count += 1
                except Exception:
                    pass
                if count > 0:
                    return
        return  # ijson 解析成功（0条也算）
    except ImportError:
        pass  # ijson 未安装，走下面的普通路径

    # ── 普通 json.load（一次性读取，处理完立即释放）────────────────
    import gc
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    chunk = (data if isinstance(data, list)
             else data.get("data", data.get("posts", data.get("items", []))))
    del data
    gc.collect()
    for post in chunk:
        if not post.get("movie"):
            post["movie"] = movie_from_file
        yield post
    del chunk
    gc.collect()


def load_sip_posts(root: Path):
    """
    只将含 SIP 编码的帖子保留在内存中，其余立即丢弃。
    返回：
        code_posts  {code: [post, ...]}   每个 SIP 编码的帖子列表
        all_sip     [post, ...]           所有 SIP 帖子（去重）
    """
    import gc
    code_posts: dict[int, list] = defaultdict(list)
    all_sip   : list[dict]      = []
    seen_ids  : set             = set()
    total_read = 0

    for path in sorted(root.rglob("*_coded.json")):
        print(f"  读取 {path.name} ...", end=" ", flush=True)
        file_sip = 0
        for p in _iter_posts_from_file(path):
            total_read += 1
            if FILTER_PATTERN in (p.get("text") or ""):
                continue
            codes = {int(c) for c in (p.get("codes") or []) if str(c).isdigit()}
            sip_hits = codes & SIP_CODE_SET
            if not sip_hits:
                continue                       # 不是 SIP 帖，直接丢弃

            p["_codes_set"] = codes
            for c in sip_hits:
                code_posts[c].append(p)

            pid = p.get("id") or id(p)
            if pid not in seen_ids:
                seen_ids.add(pid)
                all_sip.append(p)
                file_sip += 1

        print(f"SIP帖 +{file_sip}")
        gc.collect()

    print(f"共读取 {total_read:,} 条，SIP帖子（去重）{len(all_sip):,} 条")

    # 按电影+互动量排序
    for c in code_posts:
        code_posts[c].sort(key=movie_sort_key)
    all_sip.sort(key=movie_sort_key)

    return code_posts, all_sip


def movie_sort_key(p: dict) -> tuple:
    movie = normalize_movie(p.get("movie", ""))
    order = MOVIE_ORDER.index(movie) if movie in MOVIE_ORDER else 99
    return (order, -engagement(p))


HEADERS = ["序号", "电影", "用户昵称", "互动总量", "转发", "评论", "点赞",
           "发帖时间", "该帖全部编码", "帖子内容"]
COL_WIDTHS = [5, 12, 16, 10, 7, 7, 7, 17, 18, 70]


def _setup_ws(ws):
    """列宽 + 冻结首行（write-only 模式需在首次 append 前设置）。"""
    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def _header_row(ws):
    """返回带样式的表头单元格列表（WriteOnlyCell）。"""
    cells = []
    for h in HEADERS:
        c = WriteOnlyCell(ws, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = CENTER
        c.border = BDR
        cells.append(c)
    return cells


def write_sheet(ws, rows: list[dict], category: str = ""):
    """Write-only 兼容：每行 append 后立即写入磁盘，内存占用恒定。"""
    _setup_ws(ws)
    ws.append(_header_row(ws))

    for idx, p in enumerate(rows, 1):
        movie = normalize_movie(p.get("movie", ""))
        codes = p.get("_codes_set") or {int(c) for c in (p.get("codes") or []) if str(c).isdigit()}
        text  = (p.get("text") or "").replace("\n", " ").replace("\r", "")
        eng   = engagement(p)
        ws.append([
            idx, movie, get_screen_name(p),
            eng,
            int(p.get("reposts_count") or 0),
            int(p.get("comments_count") or 0),
            int(p.get("attitudes_count") or 0),
            fmt_dt(p.get("created_at", "")),
            ", ".join(str(c) for c in sorted(codes)),
            text,
        ])


def main():
    parser = argparse.ArgumentParser(description="导出SIP编码帖子明细")
    parser.add_argument("--dir",     default="data/tmp")
    parser.add_argument("--out",     default="",
                        help="输出文件名（默认自动按阈值命名）")
    parser.add_argument("--min-eng", type=int, default=0,
                        help="互动量最低阈值（转+评+赞，默认0=不过滤）")
    args = parser.parse_args()

    out_path = args.out or (
        f"sip_posts_top{args.min_eng}.xlsx" if args.min_eng > 0 else "sip_posts.xlsx"
    )

    root = Path(args.dir)
    if not root.exists():
        print(f"❌ 目录不存在：{root}")
        return

    try:
        import ijson  # noqa
        print("✔ 检测到 ijson，将使用流式解析（低内存占用）")
    except ImportError:
        print("⚠  未检测到 ijson，使用普通解析（若再次 MemoryError，请运行: pip install ijson）")

    print("扫描 SIP 编码帖子（非SIP帖子不载入内存）...")
    code_posts, all_sip = load_sip_posts(root)

    # ── 互动量过滤 ────────────────────────────────────────────────
    if args.min_eng > 0:
        print(f"过滤互动量 ≥ {args.min_eng} ...")
        code_posts = {c: [p for p in rows if engagement(p) >= args.min_eng]
                      for c, rows in code_posts.items()}
        all_sip = [p for p in all_sip if engagement(p) >= args.min_eng]
        print(f"过滤后：SIP帖子（去重）{len(all_sip):,} 条")

    # write_only=True：每行 append 立即写入临时文件，不在内存中堆积单元格对象
    wb = Workbook(write_only=True)

    def safe_sheet_name(code, name):
        s = name.replace("/", "-").replace("\\", "-").replace("*", "").replace(
            "?", "").replace("[", "").replace("]", "").replace(":", "")
        return f"c{code:02d}_{s}"[:31]

    # ── 每个编码一个 sheet ───────────────────────────────────────
    for code, category, name in SIP_CODES:
        rows = code_posts.get(code, [])
        ws = wb.create_sheet(safe_sheet_name(code, name))
        write_sheet(ws, rows, category)
        print(f"  c{code:02d} {name:<18} {len(rows):>6} 条")

    # ── 最后：所有 SIP 行为帖子汇总 ─────────────────────────────
    ws_all = wb.create_sheet("★全部SIP行为帖子")
    print(f"  ★ 写入全部SIP汇总 {len(all_sip):,} 条...", end=" ", flush=True)
    write_sheet(ws_all, all_sip)
    print("完成")

    print("保存 Excel 文件...", end=" ", flush=True)
    wb.save(out_path)
    print("完成")
    print(f"\n✅ 已导出 → {out_path}  （{len(SIP_CODES)} 个编码sheet + 1 个汇总sheet）")


if __name__ == "__main__":
    main()
