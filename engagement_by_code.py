"""
统计 coded_output.json 中各编码的转赞评总量。

用法：
    python engagement_by_code.py
    python engagement_by_code.py --input data/coded_output.json --codebook data/codebook.xlsx
"""

import argparse
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def load_code_names(codebook_path: str) -> dict[int, str]:
    wb = load_workbook(codebook_path)
    ws = wb.active
    names = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        num, name = row[0], row[1]
        if num is not None:
            names[int(num)] = str(name).strip()
    return names


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input",    default="data/coded_output.json")
    parser.add_argument("--codebook", default="data/codebook.xlsx")
    args = parser.parse_args()

    FILTER = "今年我最爱的#微博年度电影#是"
    with open(args.input, encoding="utf-8") as f:
        raw = json.load(f)
    data = [p for p in raw if FILTER not in (p.get("text") or "")]
    print(f"过滤后剩余 {len(data)} 条（去除 {len(raw)-len(data)} 条年度评选帖）\n")

    code_names = load_code_names(args.codebook)

    reposts  = defaultdict(int)
    comments = defaultdict(int)
    attitudes = defaultdict(int)
    post_counts = defaultdict(int)

    total_r = total_c = total_a = 0

    for post in data:
        r = int(post.get("reposts_count") or 0)
        c = int(post.get("comments_count") or 0)
        a = int(post.get("attitudes_count") or 0)
        total_r += r
        total_c += c
        total_a += a
        for code in post.get("codes", []):
            code = int(code)
            reposts[code]   += r
            comments[code]  += c
            attitudes[code] += a
            post_counts[code] += 1

    total_eng = total_r + total_c + total_a
    all_codes = sorted(set(code_names) | set(reposts))

    print(f"{'编码':<5} {'名称':<20} {'帖数':<7} {'转发':<10} {'评论':<10} {'点赞':<12} {'互动总量':<12} {'占总互动%'}")
    print("─" * 90)
    for code in all_codes:
        name  = code_names.get(code, "")
        n     = post_counts[code]
        r     = reposts[code]
        c     = comments[code]
        a     = attitudes[code]
        eng   = r + c + a
        pct   = eng / total_eng * 100 if total_eng else 0
        print(f"{code:<5} {name:<20} {n:<7} {r:<10,} {c:<10,} {a:<12,} {eng:<12,} {pct:.1f}%")

    print("─" * 90)
    print(f"{'全库合计':<26} {len(data):<7} {total_r:<10,} {total_c:<10,} {total_a:<12,} {total_eng:<12,} 100.0%")


if __name__ == "__main__":
    main()
