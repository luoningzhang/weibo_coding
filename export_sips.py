"""
导出高互动 SIPs 帖子（含编码17，互动量 ≥ 阈值）

用法：
    python export_sips.py                          # 默认阈值100，输出 sips_top.xlsx
    python export_sips.py --min-eng 500            # 自定义阈值
    python export_sips.py --out my_sips.xlsx
    python export_sips.py --dir data/tmp --min-eng 100 --out sips_top.xlsx
"""

import argparse
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side


SIPS_CODE = 17
FILTER_PATTERN = "今年我最爱的#微博年度电影#是"

MOVIE_ORDER = ["流浪地球2", "消失的她", "封神", "长安三万里", "孤注一掷"]


def normalize_movie(raw: str) -> str:
    return raw.removesuffix("_电影").strip()


def engagement(p: dict) -> int:
    return (int(p.get("reposts_count") or 0)
            + int(p.get("comments_count") or 0)
            + int(p.get("attitudes_count") or 0))


def load_posts(root: Path) -> list[dict]:
    posts = []
    for path in sorted(root.rglob("*_coded.json")):
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        chunk = data if isinstance(data, list) else data.get(
            "data", data.get("posts", data.get("items", [])))
        movie_from_file = path.stem.removesuffix("_coded").removesuffix("_电影")
        for p in chunk:
            if not p.get("movie"):
                p["movie"] = movie_from_file
            posts.append(p)
    return posts


def get_screen_name(p: dict) -> str:
    u = p.get("user")
    if isinstance(u, dict):
        return u.get("screen_name", "")
    return ""


def main():
    parser = argparse.ArgumentParser(description="导出高互动SIPs帖子")
    parser.add_argument("--dir",     default="data/tmp")
    parser.add_argument("--min-eng", type=int, default=100,
                        help="互动量最低阈值（转+评+赞，默认100）")
    parser.add_argument("--out",     default="sips_top.xlsx")
    args = parser.parse_args()

    root = Path(args.dir)
    if not root.exists():
        print(f"❌ 目录不存在：{root}")
        return

    posts = load_posts(root)
    print(f"已加载 {len(posts):,} 条帖子")

    # 筛选：含编码17 + 互动量 ≥ 阈值 + 排除模板帖
    sips = []
    for p in posts:
        if FILTER_PATTERN in (p.get("text") or ""):
            continue
        codes = [int(c) for c in (p.get("codes") or []) if str(c).isdigit()]
        if SIPS_CODE not in codes:
            continue
        eng = engagement(p)
        if eng >= args.min_eng:
            sips.append({**p, "_eng": eng, "_codes": codes})

    # 按电影顺序 + 互动量降序排序
    def sort_key(p):
        movie = normalize_movie(p.get("movie", ""))
        order = MOVIE_ORDER.index(movie) if movie in MOVIE_ORDER else 99
        return (order, -p["_eng"])

    sips.sort(key=sort_key)
    print(f"符合条件（含code17，互动≥{args.min_eng}）：{len(sips):,} 条")

    # ── Excel 输出 ──────────────────────────────────────────────
    wb = Workbook()

    thin  = Side(style="thin")
    BDR   = Border(left=thin, right=thin, top=thin, bottom=thin)
    HDR_FILL = PatternFill("solid", fgColor="2F5597")
    HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
    CENTER   = Alignment(horizontal="center", vertical="top", wrap_text=False)
    WRAP     = Alignment(horizontal="left",   vertical="top", wrap_text=True)

    # ── Sheet 1：帖子明细 ────────────────────────────────────────
    ws = wb.active
    ws.title = f"SIPs高互动（≥{args.min_eng}）"

    headers = ["序号", "电影", "用户昵称", "互动总量", "转发", "评论", "点赞",
               "发帖时间", "编码列表", "帖子内容（前300字）"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTER
        cell.border = BDR
    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"

    for idx, p in enumerate(sips, 1):
        movie  = normalize_movie(p.get("movie", ""))
        text   = (p.get("text") or "").replace("\n", " ").replace("\r", "")[:300]
        codes  = ", ".join(str(c) for c in sorted(p["_codes"]))
        dt_raw = p.get("created_at", "")
        try:
            dt = datetime.strptime(dt_raw, "%a %b %d %H:%M:%S +0800 %Y")
            dt_str = dt.strftime("%Y-%m-%d %H:%M")
        except Exception:
            dt_str = dt_raw

        row = [idx, movie, get_screen_name(p),
               p["_eng"],
               int(p.get("reposts_count") or 0),
               int(p.get("comments_count") or 0),
               int(p.get("attitudes_count") or 0),
               dt_str, codes, text]
        ws.append(row)
        r = ws.max_row
        for ci, cell in enumerate(ws[r], 1):
            cell.border = BDR
            cell.alignment = WRAP if ci == len(headers) else CENTER
        ws.row_dimensions[r].height = 15

    # 列宽
    col_widths = [6, 12, 16, 10, 8, 8, 8, 18, 20, 60]
    for i, w in enumerate(col_widths, 1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── Sheet 2：按电影汇总 ──────────────────────────────────────
    ws2 = wb.create_sheet("按电影汇总")
    ws2.append(["电影", f"SIPs高互动帖数（互动≥{args.min_eng}）",
                "互动总量合计", "平均互动量", "最高互动量"])
    for cell in ws2[1]:
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTER
        cell.border = BDR

    from collections import defaultdict
    mv_rows = defaultdict(list)
    for p in sips:
        mv_rows[normalize_movie(p.get("movie", ""))].append(p["_eng"])

    for movie in MOVIE_ORDER:
        engs = mv_rows.get(movie, [])
        n = len(engs)
        total = sum(engs)
        avg   = round(total / n, 1) if n else 0
        mx    = max(engs) if engs else 0
        ws2.append([movie, n, total, avg, mx])
        for cell in ws2[ws2.max_row]:
            cell.alignment = CENTER
            cell.border = BDR

    for col in ws2.columns:
        ws2.column_dimensions[col[0].column_letter].width = max(
            len(str(c.value or "")) for c in col) + 4

    wb.save(args.out)
    print(f"✅ 已导出 → {args.out}  （{len(sips)} 条帖子）")


if __name__ == "__main__":
    main()
