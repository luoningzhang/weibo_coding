"""
统计 coded_output.json 中各一阶编码的频次和占比。

用法：
    python count_codes.py
    python count_codes.py --input data/coded_output.json --codebook data/codebook.xlsx
"""

import argparse
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def load_code_names(codebook_path: str) -> dict[int, str]:
    wb = load_workbook(codebook_path)
    ws = wb.active
    names = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        num, name = row[0], row[1]
        if num is not None:
            names[int(num)] = str(name).strip()
    return names


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input",    default="data/coded_output.json")
    parser.add_argument("--codebook", default="data/codebook.xlsx")
    args = parser.parse_args()

    FILTER = "今年我最爱的#微博年度电影#是"
    with open(args.input, encoding="utf-8") as f:
        raw = json.load(f)
    data = [p for p in raw if FILTER not in (p.get("text") or "")]
    print(f"过滤后剩余 {len(data)} 条（去除 {len(raw)-len(data)} 条年度评选帖）\n")

    code_names = load_code_names(args.codebook)

    counter = Counter()
    coded_posts = 0
    for post in data:
        codes = post.get("codes", [])
        if codes:
            coded_posts += 1
        for code in codes:
            counter[int(code)] += 1

    total = len(data)
    print(f"总帖子数：{total}　　有编码：{coded_posts}　　无编码：{total - coded_posts}")
    print(f"编码实例总数：{sum(counter.values())}")
    print()
    print(f"{'编码':<5} {'名称':<20} {'频次':<8} {'占所有帖子%':<14} {'占有编码帖子%'}")
    print("─" * 65)
    for code in sorted(counter):
        cnt = counter[code]
        name = code_names.get(code, "")
        pct_all = cnt / total * 100
        pct_coded = cnt / coded_posts * 100 if coded_posts else 0
        print(f"{code:<5} {name:<20} {cnt:<8} {pct_all:<14.1f} {pct_coded:.1f}%")
    print("─" * 65)

    # 列出有定义但频次为0的编码
    missing = sorted(set(code_names) - set(counter))
    if missing:
        print(f"\n未出现的编码：{missing}")


if __name__ == "__main__":
    main()
