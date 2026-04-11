"""
各电影数据管道各环节统计

用法:
    python pipeline_stats.py                     # 扫描 data/tmp/
    python pipeline_stats.py --dir data/tmp      # 指定根目录
    python pipeline_stats.py --excel pipeline_stats.xlsx  # 导出 Excel
"""

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

# 按顺序定义各阶段：(阶段名, 文件后缀)
STAGES = [
    ("①原始",               "_电影.json"),
    ("②预清洗",              "_电影_cleaned.json"),
    ("③去认证用户",           "_电影_cleaned_no_verified.json"),
    ("④去V类",               "_电影_cleaned_no_verified_filtered.json"),
    ("⑤去S类",               "_电影_classified_F.json"),
    ("⑥编码完成",             "_电影_coded.json"),
]


def load_posts(path: Path) -> list[dict]:
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    if isinstance(data, list):
        return data
    for key in ("data", "posts", "items"):
        if key in data and isinstance(data[key], list):
            return data[key]
    return []


def get_user_id(post: dict):
    u = post.get("user")
    if isinstance(u, dict):
        return u.get("id") or u.get("user_id")
    return u or post.get("user_id")


def stage_stats(posts: list[dict]) -> tuple[int, int]:
    """返回 (条数, 不同用户数)"""
    count = len(posts)
    users = len({get_user_id(p) for p in posts if get_user_id(p) is not None})
    return count, users


def find_movies(root: Path) -> list[str]:
    """扫描子目录，返回电影名列表"""
    movies = sorted(
        d.name.replace("_电影", "")
        for d in root.iterdir()
        if d.is_dir() and not d.name.startswith(".")
    )
    return movies


def collect(root: Path, movie: str) -> list[tuple[str, int | None, int | None]]:
    """返回 [(阶段名, 条数, 用户数), ...]，文件缺失时为 None"""
    folder = root / f"{movie}_电影"
    if not folder.exists():
        folder = root / movie  # fallback：直接以电影名命名的目录
    rows = []
    for stage_name, suffix in STAGES:
        path = folder / f"{movie}{suffix}"
        if path.exists():
            posts = load_posts(path)
            cnt, users = stage_stats(posts)
            rows.append((stage_name, cnt, users))
        else:
            rows.append((stage_name, None, None))
    return rows


def print_report(root: Path, movies: list[str]):
    sep = "=" * 72
    hdr = f"  {'阶段':<12} {'条数':>8} {'不同用户':>10} {'条数保留%':>10} {'用户保留%':>10}"
    for movie in movies:
        rows = collect(root, movie)
        print(f"\n{sep}")
        print(f"  电影：{movie}")
        print(sep)
        print(hdr)
        print(f"  {'-'*68}")
        base_cnt = base_usr = None
        for stage, cnt, users in rows:
            if cnt is None:
                print(f"  {stage:<12} {'—':>8} {'—':>10} {'—':>10} {'—':>10}")
                continue
            if base_cnt is None:
                base_cnt, base_usr = cnt, users
            pct_cnt = cnt  / base_cnt * 100 if base_cnt else 0
            pct_usr = users / base_usr * 100 if base_usr else 0
            print(f"  {stage:<12} {cnt:>8,} {users:>10,} {pct_cnt:>9.1f}% {pct_usr:>9.1f}%")
    print(f"\n{sep}\n")


def export_excel(root: Path, movies: list[str], out_path: str):
    wb = Workbook()

    from openpyxl.styles import Border, Side
    from openpyxl.utils import get_column_letter

    thin      = Side(style="thin")
    BDR       = PatternFill()   # placeholder
    HDR_FILL  = PatternFill("solid", fgColor="2F5597")
    HDR_FONT  = Font(bold=True, color="FFFFFF", size=10)
    CAT_FILL  = PatternFill("solid", fgColor="D6E4F0")
    TOT_FILL  = PatternFill("solid", fgColor="FFF2CC")
    BODY_FONT = Font(size=10)
    CENTER    = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT      = Alignment(horizontal="left",   vertical="center")
    BORDER    = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_row(ws, row_i, fill=None, bold=False):
        for cell in ws[row_i]:
            cell.font   = Font(bold=bold, size=10)
            cell.border = BORDER
            cell.alignment = CENTER
            if fill:
                cell.fill = fill

    # ── 收集所有数据 ──────────────────────────────────────────────
    # data[movie] = [(stage_name, cnt, users), ...]
    all_data = {m: collect(root, m) for m in movies}

    # 阶段基数（各电影第一个非None阶段）
    base = {}
    for m in movies:
        for _, cnt, users in all_data[m]:
            if cnt is not None:
                base[m] = (cnt, users)
                break

    # ══════════════════════════════════════════════════════════════
    #  Sheet 1：宽表 — 每阶段帖子数 + 保留率
    # ══════════════════════════════════════════════════════════════
    ws = wb.active
    ws.title = "管道数据保留"

    # 表头：阶段 | 电影1 | 电影2 | ... | 全部电影合计
    hdr = ["筛选阶段"] + movies + ["全部电影\n合计"]
    ws.append(hdr)
    for cell in ws[1]:
        cell.fill      = HDR_FILL
        cell.font      = HDR_FONT
        cell.alignment = CENTER
        cell.border    = BORDER
    ws.row_dimensions[1].height = 32

    # 数据行：每阶段一行
    n_stages = len(STAGES)
    stage_totals = [0] * n_stages   # 全部电影各阶段合计

    for si, (stage_name, _) in enumerate(STAGES):
        row = [stage_name]
        row_total = 0
        for m in movies:
            cnt = all_data[m][si][1]
            row.append(cnt if cnt is not None else "—")
            if cnt is not None:
                row_total += cnt
        row.append(row_total)
        stage_totals[si] = row_total
        ws.append(row)
        style_row(ws, ws.max_row)
        ws[ws.max_row][0].alignment = LEFT

    # 空行
    ws.append([""] * (len(movies) + 2))

    # 保留率子表（相对各自①原始阶段）
    ws.append(["保留率（相对①原始）"] + [""] * (len(movies) + 1))
    ws[ws.max_row][0].font = Font(bold=True, size=10)
    ws[ws.max_row][0].fill = CAT_FILL

    for si, (stage_name, _) in enumerate(STAGES):
        row = [stage_name]
        for m in movies:
            cnt      = all_data[m][si][1]
            base_cnt = base.get(m, (None,))[0]
            if cnt is not None and base_cnt:
                row.append(f"{cnt/base_cnt*100:.1f}%")
            else:
                row.append("—")
        # 全部电影合计保留率
        tot_base = stage_totals[0] if stage_totals[0] else None
        tot_cnt  = stage_totals[si]
        row.append(f"{tot_cnt/tot_base*100:.1f}%" if tot_base else "—")
        ws.append(row)
        style_row(ws, ws.max_row)
        ws[ws.max_row][0].alignment = LEFT

    # 空行
    ws.append([""] * (len(movies) + 2))

    # ── 不同用户数子表 ──────────────────────────────────────────────
    ws.append(["不同用户数"] + [""] * (len(movies) + 1))
    ws[ws.max_row][0].font = Font(bold=True, size=10)
    ws[ws.max_row][0].fill = CAT_FILL

    user_totals = [0] * n_stages   # 全部电影各阶段用户合计

    for si, (stage_name, _) in enumerate(STAGES):
        row = [stage_name]
        row_total = 0
        for m in movies:
            users = all_data[m][si][2]
            row.append(users if users is not None else "—")
            if users is not None:
                row_total += users
        row.append(row_total)
        user_totals[si] = row_total
        ws.append(row)
        style_row(ws, ws.max_row)
        ws[ws.max_row][0].alignment = LEFT

    # 空行
    ws.append([""] * (len(movies) + 2))

    # ── 用户保留率子表（相对各自①原始阶段）───────────────────────────
    ws.append(["用户保留率（相对①原始）"] + [""] * (len(movies) + 1))
    ws[ws.max_row][0].font = Font(bold=True, size=10)
    ws[ws.max_row][0].fill = CAT_FILL

    for si, (stage_name, _) in enumerate(STAGES):
        row = [stage_name]
        for m in movies:
            users      = all_data[m][si][2]
            base_users = base.get(m, (None, None))[1]
            if users is not None and base_users:
                row.append(f"{users/base_users*100:.1f}%")
            else:
                row.append("—")
        # 全部电影合计用户保留率
        tot_base_u = user_totals[0] if user_totals[0] else None
        tot_u      = user_totals[si]
        row.append(f"{tot_u/tot_base_u*100:.1f}%" if tot_base_u else "—")
        ws.append(row)
        style_row(ws, ws.max_row)
        ws[ws.max_row][0].alignment = LEFT

    # 列宽
    ws.column_dimensions["A"].width = 22
    for ci in range(2, len(movies) + 3):
        ws.column_dimensions[get_column_letter(ci)].width = 14

    wb.save(out_path)
    print(f"✅ Excel 已导出 → {out_path}")


def main():
    parser = argparse.ArgumentParser(description="数据管道各环节条数与用户数统计")
    parser.add_argument("--dir",   default="data/tmp", help="数据根目录")
    parser.add_argument("--excel", default="",         help="导出 Excel 路径")
    args = parser.parse_args()

    root = Path(args.dir)
    if not root.exists():
        print(f"❌ 目录不存在：{root}")
        return

    movies = find_movies(root)
    if not movies:
        print(f"❌ {root} 下未找到任何子目录")
        return

    print(f"找到 {len(movies)} 部电影：{', '.join(movies)}")
    print_report(root, movies)

    if args.excel:
        export_excel(root, movies, args.excel)


if __name__ == "__main__":
    main()
