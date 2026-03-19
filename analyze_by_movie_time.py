"""
按电影、上映阶段分析所有编码的分布与互动量。
输出三张 CSV + 控制台摘要。

上映日期（院线首映）：
  流浪地球2   2023-01-22
  消失的她    2023-06-22
  封神第一部  2023-07-20
  长安三万里  2023-07-08
  孤注一掷    2023-08-08

时间分段（HOT_DAYS 可调）：
  上映前   created_at < 首映日
  上映期   首映日 ≤ created_at ≤ 首映日 + HOT_DAYS天
  长尾期   created_at > 首映日 + HOT_DAYS天

用法：
    python analyze_by_movie_time.py
    python analyze_by_movie_time.py --input data/coded_output.json --out-dir data
"""

import argparse
import csv
import json
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

RELEASE_DATES = {
    "流浪地球2_电影": datetime(2023, 1, 22),
    "消失的她_电影":  datetime(2023, 6, 22),
    "封神_电影":      datetime(2023, 7, 20),
    "长安三万里_电影": datetime(2023, 7, 8),
    "孤注一掷_电影":  datetime(2023, 8, 8),
}
HOT_DAYS = 30
PHASES = ["上映前", "上映期", "长尾期"]
FILTER_PATTERN = "今年我最爱的#微博年度电影#是"



def load_code_names(codebook_path: str) -> dict[int, str]:
    wb = load_workbook(codebook_path)
    ws = wb.active
    names = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        num, name = row[0], row[1]
        if num is not None:
            names[int(num)] = str(name).strip()
    return names


def parse_dt(s: str):
    try:
        return datetime.strptime(s, "%a %b %d %H:%M:%S +0800 %Y")
    except Exception:
        return None


def get_phase(dt: datetime, release: datetime) -> str:
    if dt < release:
        return "上映前"
    elif dt <= release + timedelta(days=HOT_DAYS):
        return "上映期"
    else:
        return "长尾期"


def eng(post: dict) -> int:
    return (int(post.get("reposts_count") or 0)
            + int(post.get("comments_count") or 0)
            + int(post.get("attitudes_count") or 0))


def sep(char="─", width=72):
    print(char * width)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input",    default="data/coded_output.json")
    parser.add_argument("--codebook", default="data/codebook.xlsx")
    parser.add_argument("--out-dir",  default="data")
    args = parser.parse_args()

    with open(args.input, encoding="utf-8") as f:
        data = json.load(f)

    code_names = load_code_names(args.codebook)
    all_codes  = sorted(code_names)
    out_dir    = Path(args.out_dir)

    # ── 累加器 ────────────────────────────────────────────────────
    # [movie][code] -> {posts, reposts, comments, attitudes}
    movie_code = defaultdict(lambda: defaultdict(lambda: [0, 0, 0, 0]))
    # [movie][phase][code] -> {posts, reposts, comments, attitudes}
    mp_code    = defaultdict(lambda: defaultdict(lambda: defaultdict(lambda: [0, 0, 0, 0])))
    # [movie][phase] -> {total_posts, total_eng}
    mp_totals  = defaultdict(lambda: defaultdict(lambda: [0, 0]))
    # 全库合计
    total_posts_all = 0
    total_eng_all   = 0

    filtered = [p for p in data if FILTER_PATTERN not in (p.get("text") or "")]
    n_dropped = len(data) - len(filtered)
    print(f'过滤含"{FILTER_PATTERN}"的帖子：{n_dropped} 条，剩余 {len(filtered)} 条\n')

    for post in filtered:
        movie   = post.get("movie", "")
        dt      = parse_dt(post.get("created_at", ""))
        r       = int(post.get("reposts_count") or 0)
        c       = int(post.get("comments_count") or 0)
        a       = int(post.get("attitudes_count") or 0)
        e       = r + c + a
        codes   = [int(x) for x in post.get("codes", [])]
        release = RELEASE_DATES.get(movie)
        total_posts_all += 1
        total_eng_all   += e

        for code in codes:
            movie_code[movie][code][0] += 1
            movie_code[movie][code][1] += r
            movie_code[movie][code][2] += c
            movie_code[movie][code][3] += a

        if dt and release:
            phase = get_phase(dt, release)
            mp_totals[movie][phase][0] += 1
            mp_totals[movie][phase][1] += e
            for code in codes:
                mp_code[movie][phase][code][0] += 1
                mp_code[movie][phase][code][1] += r
                mp_code[movie][phase][code][2] += c
                mp_code[movie][phase][code][3] += a

    # ══════════════════════════════════════════════════════════════
    #  CSV 1：电影 × 编码（全编码）
    # ══════════════════════════════════════════════════════════════
    csv1 = out_dir / "by_movie_code.csv"
    with open(csv1, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["电影", "编码", "编码名", "帖数", "转发", "评论", "点赞", "互动总量"])
        for movie in sorted(RELEASE_DATES):
            for code in all_codes:
                v = movie_code[movie][code]
                w.writerow([movie, code, code_names.get(code, ""),
                             v[0], v[1], v[2], v[3], v[1]+v[2]+v[3]])

    # ══════════════════════════════════════════════════════════════
    #  CSV 2：电影 × 阶段 × 编码（全编码）
    # ══════════════════════════════════════════════════════════════
    csv2 = out_dir / "by_movie_phase_code.csv"
    with open(csv2, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["电影", "阶段", "编码", "编码名", "帖数", "转发", "评论", "点赞", "互动总量"])
        for movie in sorted(RELEASE_DATES):
            for phase in PHASES:
                for code in all_codes:
                    v = mp_code[movie][phase][code]
                    w.writerow([movie, phase, code, code_names.get(code, ""),
                                 v[0], v[1], v[2], v[3], v[1]+v[2]+v[3]])

    # ══════════════════════════════════════════════════════════════
    #  CSV 3：电影 × 阶段 帖数/互动汇总
    # ══════════════════════════════════════════════════════════════
    csv3 = out_dir / "by_movie_phase_totals.csv"
    with open(csv3, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["电影", "首映日", "阶段", "帖数", "互动总量"])
        for movie in sorted(RELEASE_DATES):
            rel = RELEASE_DATES[movie].strftime("%Y-%m-%d")
            for phase in PHASES:
                v = mp_totals[movie][phase]
                w.writerow([movie, rel, phase, v[0], v[1]])

    print(f"已输出 CSV：\n  {csv1}\n  {csv2}\n  {csv3}")

    # ══════════════════════════════════════════════════════════════
    #  控制台：【一】五部电影 × 全编码
    # ══════════════════════════════════════════════════════════════
    print("\n" + "="*72)
    print("【一】五部电影 × 全编码（帖数 / 互动总量）")
    for movie in sorted(RELEASE_DATES):
        total_n = sum(movie_code[movie][c][0] for c in all_codes)
        total_e = sum(movie_code[movie][c][1]+movie_code[movie][c][2]+movie_code[movie][c][3]
                      for c in all_codes)
        sep()
        print(f"  {movie}  |  帖子总数：{movie_posts_n(filtered, movie):,}  |  编码实例：{total_n:,}  |  互动总量：{total_e:,}")
        sep("·")
        print(f"  {'编码':<5} {'名称':<18} {'帖数':>7} {'互动总量':>12}")
        sep("·")
        for code in all_codes:
            v = movie_code[movie][code]
            n, e = v[0], v[1]+v[2]+v[3]
            if n == 0:
                continue
            print(f"  {code:<5} {code_names.get(code,''):<18} {n:>7,} {e:>12,}")

    # ══════════════════════════════════════════════════════════════
    #  控制台：【二】每部电影三阶段 × 全编码
    # ══════════════════════════════════════════════════════════════
    print("\n" + "="*72)
    print("【二】每部电影 × 上映阶段 × 全编码")
    for movie in sorted(RELEASE_DATES):
        rel = RELEASE_DATES[movie].strftime("%Y-%m-%d")
        print(f"\n  ════ {movie}  首映：{rel} ════")
        for phase in PHASES:
            pt, pe = mp_totals[movie][phase]
            print(f"\n  ── {phase}  帖数：{pt:,}  互动：{pe:,} ──")
            print(f"  {'编码':<5} {'名称':<18} {'帖数':>7} {'互动总量':>12}")
            sep("·")
            any_row = False
            for code in all_codes:
                v = mp_code[movie][phase][code]
                n, e = v[0], v[1]+v[2]+v[3]
                if n == 0:
                    continue
                print(f"  {code:<5} {code_names.get(code,''):<18} {n:>7,} {e:>12,}")
                any_row = True
            if not any_row:
                print("  （无数据）")


def movie_posts_n(data, movie):
    return sum(1 for p in data if p.get("movie") == movie)


if __name__ == "__main__":
    main()
